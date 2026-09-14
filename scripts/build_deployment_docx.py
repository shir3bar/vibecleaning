"""Build IT-distribution DOCX copies from the authoritative Markdown files."""

from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCES = (
    ROOT / "docs" / "IT_DEPLOYMENT_REQUIREMENTS.md",
    ROOT / "docs" / "WINDOWS_COMPATIBILITY_TODO.md",
)


def _add_code(document: Document, lines: list[str]) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def build(source: Path) -> Path:
    document = Document()
    document.core_properties.title = source.stem.replace("_", " ").title()
    document.core_properties.subject = "Generated from the matching Markdown source"
    code_lines: list[str] | None = None
    paragraph_lines: list[str] = []
    list_paragraph = None

    def flush_paragraph() -> None:
        nonlocal paragraph_lines
        if paragraph_lines:
            document.add_paragraph(" ".join(paragraph_lines).replace("`", ""))
            paragraph_lines = []

    for raw_line in source.read_text(encoding="utf-8").splitlines():
        line = raw_line.rstrip()
        if line.startswith("```"):
            if code_lines is None:
                flush_paragraph()
                code_lines = []
            else:
                _add_code(document, code_lines)
                code_lines = None
            list_paragraph = None
            continue
        if code_lines is not None:
            code_lines.append(line)
            continue
        if not line:
            flush_paragraph()
            list_paragraph = None
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            flush_paragraph()
            document.add_heading(heading.group(2), level=len(heading.group(1)))
            list_paragraph = None
            continue
        numbered = re.match(r"^\d+\.\s+(.+)$", line)
        if numbered:
            flush_paragraph()
            list_paragraph = document.add_paragraph(
                numbered.group(1).replace("`", ""), style="List Number"
            )
            continue
        bullet = re.match(r"^-\s+(?:\[([ xX])\]\s+)?(.+)$", line)
        if bullet:
            flush_paragraph()
            checkbox = "☒ " if str(bullet.group(1) or "").lower() == "x" else (
                "☐ " if bullet.group(1) is not None else ""
            )
            list_paragraph = document.add_paragraph(
                checkbox + bullet.group(2).replace("`", ""), style="List Bullet"
            )
            continue
        if list_paragraph is not None and raw_line[:1].isspace():
            list_paragraph.add_run(" " + line.strip().replace("`", ""))
            continue
        list_paragraph = None
        paragraph_lines.append(line.strip())
    flush_paragraph()
    if code_lines is not None:
        _add_code(document, code_lines)
    destination = source.with_suffix(".docx")
    document.save(destination)
    return destination


def main() -> None:
    for source in SOURCES:
        print(build(source).relative_to(ROOT))


if __name__ == "__main__":
    main()
